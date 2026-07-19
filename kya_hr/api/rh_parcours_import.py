# Copyright (c) 2026, KYA-Energy Group
"""Import du classeur RH : registre du personnel + parcours de carrière.

La RH tient son personnel dans « KYA-Energy_RH_Base_de_donnees_et_Tableau_de_
bord.xlsx » : une feuille `Personnel` (une ligne par salarié) et une feuille
`Evolution carriere` (une ligne par évènement). Ce module verse ces deux
feuilles dans `Salarie KYA` et `Evolution Carriere KYA`.

Deux principes, identiques à l'import du stock :

* **On ne contredit jamais une saisie faite dans l'ERP.** À la création on
  renseigne tout ce que le classeur sait ; ensuite on ne remplit que les
  champs restés VIDES.
* **Rejouable sans dégât.** Un évènement est identifié par (salarié, date de
  début, nature) : réimporter le même classeur ne crée pas de doublon.

Le classeur ne datant pas tous les évènements, une **embauche** est déduite de
la date d'embauche du registre pour les salariés qui n'ont aucun évènement :
sans ce point de départ leur parcours resterait vide à l'écran alors que
l'information existe.
"""
import base64
import io

import frappe
from frappe import _
from frappe.utils import flt, getdate

FEUILLE_PERSONNEL = "Personnel"
FEUILLE_CARRIERE = "Evolution carriere"

_ROLES = {"System Manager", "Responsable RH", "HR Manager", "Assistant(e) RH"}

# Le classeur emploie un vocabulaire plus riche que le DocType. On rattache
# chaque libellé à la nature existante la plus proche et on conserve le libellé
# d'origine dans l'observation : aucune information de la RH n'est perdue.
_NATURES = {
    "embauche": "Embauche",
    "recrutement": "Embauche",
    "promotion": "Promotion",
    "mutation": "Mutation",
    "affectation": "Mutation",
    "renouvellement": "Renouvellement",
    "augmentation salariale": "Augmentation salariale",
    "augmentation": "Augmentation salariale",
    "retrogradation": "Rétrogradation",
    "rétrogradation": "Rétrogradation",
    "fin de contrat": "Fin de contrat",
    "depart": "Fin de contrat",
    "départ": "Fin de contrat",
    # « Evolution du statut contractuel en CDI » : un changement de contrat,
    # que le DocType nomme Renouvellement.
    "statut contractuel": "Renouvellement",
    "titularisation": "Renouvellement",
}
_NATURES_CONNUES = {
    "Embauche", "Promotion", "Mutation", "Renouvellement",
    "Augmentation salariale", "Rétrogradation", "Fin de contrat",
}

# Marqueur de simulation : salarié qui SERAIT créé par un import réel. Il doit
# rester DISTINCT par matricule, sinon tous les futurs salariés se confondent en
# une seule identité et l'embauche déduite n'est comptée qu'une seule fois.
_A_CREER = "\x00a-creer:"


def _a_creer(matricule):
    return _A_CREER + str(matricule)


def _est_a_creer(salarie):
    return bool(salarie) and str(salarie).startswith(_A_CREER)

_CONTRATS = {"cdi": "CDI", "cdd": "CDD", "stage": "Stagiaire",
             "stagiaire": "Stagiaire", "prestataire": "Prestataire",
             "interim": "Intérim", "intérim": "Intérim"}


def _txt(v):
    if v is None:
        return ""
    return " ".join(str(v).split()).strip()


def _date(v):
    """Date exploitable, ou None. Le classeur contient des cellules d'erreur
    Excel (#VALUE!, #N/A) qu'il ne faut surtout pas prendre pour des dates."""
    if v is None:
        return None
    s = _txt(v)
    if not s or s.startswith("#"):
        return None
    try:
        return getdate(v if hasattr(v, "year") else s)
    except Exception:
        return None


def _nature(libelle):
    """(nature retenue, libellé d'origine si différent)."""
    brut = _txt(libelle)
    if not brut:
        return None, ""
    if brut in _NATURES_CONNUES:
        return brut, ""
    cle = brut.lower()
    if cle in _NATURES:
        return _NATURES[cle], brut
    # « Evolution du statut contractuel en CDI », « Affectation intérieur… » :
    # on reconnaît le premier mot-clé rencontré plutôt que d'abandonner.
    for motif, cible in _NATURES.items():
        if motif in cle:
            return cible, brut
    return None, brut


def _departement(code):
    """Rattache « DST », « DSS »… à un `Departement KYA`. Rien de sûr → vide,
    plutôt que de ranger la personne dans le mauvais service."""
    code = _txt(code)
    if not code:
        return None
    for champ in ("name", "nom_departement", "abreviation"):
        try:
            trouve = frappe.db.get_value("Departement KYA", {champ: code}, "name")
            if trouve:
                return trouve
        except Exception:
            pass
    return None


def _lire_feuille(wb, nom):
    """Lignes d'une feuille sous forme de dictionnaires. La ligne d'en-tête
    n'est pas la première : les classeurs KYA ont un titre et un sous-titre."""
    if nom not in wb.sheetnames:
        return []
    ws = wb[nom]
    lignes = list(ws.iter_rows(values_only=True))
    entete, debut = None, 0
    for i, r in enumerate(lignes[:12]):
        cellules = [_txt(c).lower() for c in r]
        if "matricule" in cellules or "mle" in cellules:
            entete, debut = [_txt(c) for c in r], i + 1
            break
    if entete is None:
        return []
    out = []
    for r in lignes[debut:]:
        d = {}
        for i, cle in enumerate(entete):
            if cle and i < len(r):
                d[cle.lower()] = r[i]
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out


def _col(d, *noms):
    """Première colonne présente parmi `noms` (comparaison souple)."""
    for n in noms:
        n = n.lower()
        for cle, val in d.items():
            if cle == n or cle.startswith(n):
                return val
    return None


# ── Registre du personnel ────────────────────────────────────────────────
def _valeurs_salarie(d):
    nom, prenoms = _txt(_col(d, "nom")), _txt(_col(d, "prénoms", "prenoms", "prénom"))
    contrat = _txt(_col(d, "type de contrat"))
    vals = {
        "nom": nom,
        "prenoms": prenoms,
        "nom_complet": " ".join(x for x in (nom, prenoms) if x),
        "sexe": _txt(_col(d, "sexe")) or None,
        "numero_assurance_sociale": _txt(_col(d, "n° assurance", "n assurance")),
        "nationalite": _txt(_col(d, "nationalité", "nationalite")) or None,
        "statut_matrimonial": _txt(_col(d, "statut matrimonial")) or None,
        "date_naissance": _date(_col(d, "date de naissance")),
        "poste_occupe": _txt(_col(d, "poste occupé", "poste occupe")),
        "departement": _departement(_col(d, "département de r", "departement de r",
                                         "département", "departement")),
        "qualification": _txt(_col(d, "qualification")),
        "date_embauche": _date(_col(d, "date d'embauche", "date d embauche")),
        "date_debauchage": _date(_col(d, "date de débauchage", "date de debauchage")),
        "type_contrat": _CONTRATS.get(contrat.lower(), contrat or None),
    }
    enfants = _col(d, "nombre d'enfants", "nombre d enfants")
    if enfants not in (None, ""):
        vals["nombre_enfants"] = int(flt(enfants))
    vals["statut_emploi"] = "Sorti" if vals.get("date_debauchage") else "Actif"
    return {k: v for k, v in vals.items() if v not in (None, "")}


def _upsert_salarie(matricule, vals):
    """Retourne (nom du document, 'cree'|'complete'|'inchange')."""
    existant = frappe.db.get_value("Salarie KYA", {"matricule": matricule}, "name")
    if not existant:
        doc = frappe.get_doc(dict(doctype="Salarie KYA", matricule=matricule, **vals))
        doc.flags.ignore_permissions = True
        doc.insert()
        return doc.name, "cree"

    sal = frappe.get_doc("Salarie KYA", existant)
    change = False
    for champ, valeur in vals.items():
        if sal.get(champ) in (None, "", 0) and sal.get(champ) != valeur:
            sal.set(champ, valeur)
            change = True
    if not change:
        return existant, "inchange"
    sal.flags.ignore_permissions = True
    sal.save()
    return existant, "complete"


# ── Parcours de carrière ─────────────────────────────────────────────────
def _evenement_existe(salarie, date_debut, nature):
    return frappe.db.exists("Evolution Carriere KYA", {
        "salarie": salarie, "date_debut": date_debut, "nature_evolution": nature})


def _creer_evenement(salarie, date_debut, nature, **champs):
    if _evenement_existe(salarie, date_debut, nature):
        return False
    doc = frappe.get_doc(dict(
        doctype="Evolution Carriere KYA", salarie=salarie,
        date_debut=date_debut, nature_evolution=nature,
        **{k: v for k, v in champs.items() if v not in (None, "")}))
    doc.flags.ignore_permissions = True
    doc.insert()
    return True


def _valeurs_evenement(d):
    contrat = _txt(_col(d, "type de contrat"))
    renouv = _col(d, "nombre de renouvellement")
    return {
        "date_fin": _date(_col(d, "date de fin")),
        "type_contrat": _CONTRATS.get(contrat.lower(), contrat or None),
        "departement": _departement(_col(d, "département", "departement")),
        "poste_occupe": _txt(_col(d, "poste occupé", "poste occupe")),
        "categorie": _txt(_col(d, "catégorie", "categorie")) or None,
        "classe_echelon": _txt(_col(d, "classe")),
        "salaire_base": flt(_col(d, "salaire de base")) or None,
        "motif_evolution": _txt(_col(d, "motif")),
        "nombre_renouvellements": int(flt(renouv)) if renouv not in (None, "") else None,
    }


@frappe.whitelist()
def importer_classeur_rh(content_base64, filename=None, dry_run=1):
    """Verse le classeur RH dans le registre et le parcours de carrière.

    `dry_run=1` (défaut) : compte ce qui serait fait, sans rien écrire.
    """
    if not (_ROLES & set(frappe.get_roles(frappe.session.user))):
        frappe.throw(_("Réservé aux Ressources Humaines."), frappe.PermissionError)

    dry_run = int(dry_run or 0)
    import openpyxl

    raw = base64.b64decode(content_base64)
    if not (filename or "").lower().endswith((".xlsx", ".xlsm")) and raw[:2] != b"PK":
        frappe.throw(_("Le classeur RH doit être un fichier Excel (.xlsx)."))
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)

    res = {"dry_run": bool(dry_run), "feuilles": wb.sheetnames,
           "salaries_crees": 0, "salaries_completes": 0, "salaries_inchanges": 0,
           "evenements_crees": 0, "evenements_existants": 0,
           "embauches_deduites": 0, "ignorees": [], "erreurs": []}

    personnel = _lire_feuille(wb, FEUILLE_PERSONNEL)
    carriere = _lire_feuille(wb, FEUILLE_CARRIERE)
    if not personnel and not carriere:
        frappe.throw(_("Feuilles « {0} » et « {1} » introuvables ou vides dans ce classeur.")
                     .format(FEUILLE_PERSONNEL, FEUILLE_CARRIERE))

    # 1. Registre du personnel — il doit exister avant le parcours, qui s'y rattache.
    par_matricule = {}
    for d in personnel:
        matricule = _txt(_col(d, "mle", "matricule"))
        if not matricule:
            continue
        vals = _valeurs_salarie(d)
        if not vals.get("nom"):
            res["ignorees"].append("Personnel %s : nom absent" % matricule)
            continue
        try:
            if dry_run:
                existe = frappe.db.get_value("Salarie KYA", {"matricule": matricule}, "name")
                if not existe:
                    res["salaries_crees"] += 1
                else:
                    # « Complété » seulement si un champ vide serait rempli,
                    # sinon la simulation annoncerait 56 mises à jour là où
                    # l'import réel ne touchera rien.
                    actuel = frappe.db.get_value("Salarie KYA", existe,
                                                 list(vals.keys()), as_dict=True) or {}
                    comble = any(actuel.get(c) in (None, "", 0) and actuel.get(c) != v
                                 for c, v in vals.items())
                    res["salaries_completes" if comble else "salaries_inchanges"] += 1
                # En simulation le salarié n'est pas encore créé : sans ce
                # marqueur, tout son parcours serait annoncé « salarié inconnu »
                # alors que l'import réel le rattacherait sans difficulté.
                par_matricule[matricule] = (existe or _a_creer(matricule), vals)
            else:
                nom_doc, action = _upsert_salarie(matricule, vals)
                res["salaries_" + {"cree": "crees", "complete": "completes",
                                   "inchange": "inchanges"}[action]] += 1
                par_matricule[matricule] = (nom_doc, vals)
        except Exception:
            res["erreurs"].append("Personnel %s" % matricule)
            frappe.log_error(frappe.get_traceback(), "rh_parcours_import.personnel")

    # 2. Parcours de carrière.
    avec_evenement = set()
    for d in carriere:
        matricule = _txt(_col(d, "matricule", "mle"))
        if not matricule:
            continue
        salarie = (par_matricule.get(matricule) or (None,))[0] \
            or frappe.db.get_value("Salarie KYA", {"matricule": matricule}, "name")
        date_debut = _date(_col(d, "date de début", "date de debut"))
        nature, libelle = _nature(_col(d, "nature de l'évolution", "nature de l",
                                       "type d'événement", "type d"))

        # Une ligne sans date OU sans nature n'est pas un évènement : le
        # classeur pré-remplit matricule et nom pour TOUT le personnel.
        if not date_debut or not nature:
            if date_debut or libelle:
                if not date_debut:
                    motif = "date de début absente"
                elif not libelle:
                    motif = "nature de l'évolution absente"
                else:
                    motif = "nature « %s » non reconnue" % libelle
                res["ignorees"].append("Carrière %s : %s" % (matricule, motif))
            continue
        if not salarie:
            res["ignorees"].append("Carrière %s : salarié inconnu dans le registre" % matricule)
            continue

        avec_evenement.add(salarie)
        try:
            if dry_run:
                if not _est_a_creer(salarie) and _evenement_existe(salarie, date_debut, nature):
                    res["evenements_existants"] += 1
                else:
                    res["evenements_crees"] += 1
                continue
            champs = _valeurs_evenement(d)
            if libelle:
                obs = _txt(_col(d, "observation"))
                champs["observation"] = ("%s — libellé d'origine : « %s »"
                                         % (obs, libelle) if obs
                                         else "Libellé d'origine : « %s »" % libelle)
            if _creer_evenement(salarie, date_debut, nature, **champs):
                res["evenements_crees"] += 1
            else:
                res["evenements_existants"] += 1
        except Exception:
            res["erreurs"].append("Carrière %s" % matricule)
            frappe.log_error(frappe.get_traceback(), "rh_parcours_import.carriere")

    # 3. Embauche déduite : sans elle, le parcours de la plupart des salariés
    #    resterait vide alors que la date d'embauche est connue du registre.
    for matricule, (salarie, vals) in par_matricule.items():
        date_embauche = vals.get("date_embauche")
        if not salarie or not date_embauche or salarie in avec_evenement:
            continue
        if dry_run:
            if _est_a_creer(salarie) or not _evenement_existe(salarie, date_embauche, "Embauche"):
                res["embauches_deduites"] += 1
            continue
        if _evenement_existe(salarie, date_embauche, "Embauche"):
            continue
        try:
            if _creer_evenement(
                    salarie, date_embauche, "Embauche",
                    type_contrat=vals.get("type_contrat"),
                    departement=vals.get("departement"),
                    poste_occupe=vals.get("poste_occupe"),
                    motif_evolution="Recrutement initial",
                    observation="Déduite de la date d'embauche du registre du personnel."):
                res["embauches_deduites"] += 1
        except Exception:
            res["erreurs"].append("Embauche déduite %s" % matricule)
            frappe.log_error(frappe.get_traceback(), "rh_parcours_import.embauche")

    if not dry_run:
        frappe.db.commit()
    return res
