"""API du module RH « Effectifs » — indicateurs & graphiques du tableau de bord,
calculés EN DIRECT depuis le registre `Salarie KYA` (aucune valeur figée, aucun
recours à la paie ERPNext). Reproduit fidèlement la maquette du classeur RH.
"""
from __future__ import annotations

import frappe
from frappe import _
from frappe.utils import add_months, flt, get_first_day, getdate, today

_RH_ROLES = {"System Manager", "Responsable RH", "HR Manager", "HR User",
             "Assistant(e) RH", "Directeur Général", "DG", "DGA"}

# Ordre d'affichage fixe (maquette du classeur).
_NIVEAUX = ["Sans diplôme", "CEPD/BEPC", "CAP/BEP", "Baccalauréat", "Bac",
            "BAC+2/BTS/DUT", "Licence", "Master", "Ingénieur", "Doctorat", "Postdoc"]
_CONTRATS = ["CDI", "CDD", "Stagiaire", "Prestataire", "Intérim"]
_CATEGORIES = ["C", "AM", "AE"]
_TRAVAIL = ["Présentiel", "Télétravail", "Hybride"]
_TRANCHES = ["< 25 ans", "25 - 34 ans", "35 - 44 ans", "45 - 54 ans", "55 ans et +"]


def _guard():
    if not _RH_ROLES.intersection(set(frappe.get_roles(frappe.session.user))):
        frappe.throw(_("Accès réservé aux Ressources Humaines."), frappe.PermissionError)


def _age(dn, ref=None):
    if not dn:
        return None
    ref = getdate(ref or today())
    dn = getdate(dn)
    return int((ref - dn).days // 365.25)


def _anciennete(de, fin=None):
    if not de:
        return 0.0
    fin = getdate(fin or today())
    return round((fin - getdate(de)).days / 365.25, 1)


def _statut(row):
    if not row.get("date_debauchage"):
        return "Actif"
    return "Retraité" if (row.get("motif_debauchage") or "").strip().lower() == "retraite" else "Sorti"


def _tranche(age):
    if age is None:
        return None
    if age < 25:
        return _TRANCHES[0]
    if age < 35:
        return _TRANCHES[1]
    if age < 45:
        return _TRANCHES[2]
    if age < 55:
        return _TRANCHES[3]
    return _TRANCHES[4]


def _median(vals):
    vals = sorted(v for v in vals if v is not None)
    n = len(vals)
    if not n:
        return 0
    mid = n // 2
    return round(vals[mid] if n % 2 else (vals[mid - 1] + vals[mid]) / 2, 1)


@frappe.whitelist()
def dashboard_data(departement=None, periode_debut=None, periode_fin=None):
    """Retourne KPI + toutes les distributions du tableau de bord RH, en direct.

    Filtres : `departement` restreint la population ; `periode_debut/fin` borne
    les courbes d'évolution (effectif & masse salariale dans le temps)."""
    _guard()
    filters = {}
    if departement:
        filters["departement"] = departement

    rows = frappe.get_all(
        "Salarie KYA", filters=filters,
        fields=["name", "matricule", "nom_complet", "sexe", "categorie",
                "type_contrat", "type_travail", "departement", "niveau_etudes",
                "poste_occupe", "salaire_base", "date_naissance", "date_embauche",
                "date_debauchage", "motif_debauchage"],
        limit_page_length=0)

    # ── Population active (les indicateurs « photo » portent sur les actifs) ──
    actifs = [r for r in rows if _statut(r) == "Actif"]
    ages = [_age(r.date_naissance) for r in actifs]
    ages = [a for a in ages if a is not None]
    anciennetes = [_anciennete(r.date_embauche, r.date_debauchage) for r in actifs]

    def _count_by(seq, key, order=None):
        d = {}
        for r in seq:
            v = (r.get(key) or "—")
            d[v] = d.get(v, 0) + 1
        if order:
            return {"labels": order, "data": [d.get(k, 0) for k in order]}
        items = sorted(d.items(), key=lambda kv: -kv[1])
        return {"labels": [k for k, _v in items], "data": [v for _k, v in items]}

    masse = sum(flt(r.salaire_base) for r in actifs)

    kpi = {
        "effectif_total": len(actifs),
        "hommes": sum(1 for r in actifs if r.sexe == "M"),
        "femmes": sum(1 for r in actifs if r.sexe == "F"),
        "cdi": sum(1 for r in actifs if r.type_contrat == "CDI"),
        "temporaires": sum(1 for r in actifs if r.type_contrat and r.type_contrat != "CDI"),
        "masse_salariale": round(masse, 0),
        "cadres": sum(1 for r in actifs if r.categorie == "C"),
        "maitrise": sum(1 for r in actifs if r.categorie == "AM"),
        "execution": sum(1 for r in actifs if r.categorie == "AE"),
        "age_moyen": round(sum(ages) / len(ages), 1) if ages else 0,
        "age_median": _median(ages),
        "anciennete_moyenne": round(sum(anciennetes) / len(anciennetes), 1) if anciennetes else 0,
        # bonus (parcours / retraités)
        "retraites": sum(1 for r in rows if _statut(r) == "Retraité"),
        "sortis": sum(1 for r in rows if _statut(r) == "Sorti"),
        # comme le classeur (ligne « STAGIAIRES EN COURS / PRESTATAIRES EN COURS »)
        "stagiaires_en_cours": frappe.db.count("Stagiaire RH KYA", {"statut": "En cours"})
        if frappe.db.exists("DocType", "Stagiaire RH KYA") else 0,
        "prestataires_en_cours": frappe.db.count("Prestataire KYA", {"statut": "En cours"})
        if frappe.db.exists("DocType", "Prestataire KYA") else 0,
    }

    # ── Distributions (camemberts + barres) ──
    par_sexe = {"labels": ["Hommes", "Femmes"],
                "data": [kpi["hommes"], kpi["femmes"]]}
    par_categorie = _count_by(actifs, "categorie", _CATEGORIES)
    par_contrat = _count_by(actifs, "type_contrat", _CONTRATS)
    par_travail = _count_by(actifs, "type_travail", _TRAVAIL)
    par_departement = _count_by(actifs, "departement")
    par_niveau = _count_by(actifs, "niveau_etudes", _NIVEAUX)
    postes = _count_by(actifs, "poste_occupe")
    par_poste = {"labels": postes["labels"][:10], "data": postes["data"][:10]}

    tr = {t: 0 for t in _TRANCHES}
    for r in actifs:
        t = _tranche(_age(r.date_naissance))
        if t:
            tr[t] += 1
    par_tranche = {"labels": _TRANCHES, "data": [tr[t] for t in _TRANCHES]}

    # ── Évolution effectif & masse salariale dans le temps (mois par mois) ──
    evol = _evolution(rows, periode_debut, periode_fin)

    return {
        "date_str": frappe.utils.formatdate(today(), "EEEE d MMMM y"),
        "kpi": kpi,
        "par_sexe": par_sexe, "par_categorie": par_categorie,
        "par_contrat": par_contrat, "par_travail": par_travail,
        "par_departement": par_departement, "par_niveau": par_niveau,
        "par_poste": par_poste, "par_tranche": par_tranche,
        "evolution": evol,
        "departements": frappe.get_all("Departement KYA", filters={"actif": 1},
                                       fields=["code", "libelle"], order_by="code"),
    }


def _norm(s):
    """Minuscule + sans accents + apostrophes→espaces, pour comparer des en-têtes."""
    import unicodedata
    s = str(s or "").replace("'", " ").replace("’", " ").replace("`", " ")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


# En-tête (normalisé, sous-chaîne) -> champ Salarie KYA.
_IMPORT_MAP = [
    ("mle", "matricule"), ("assurance sociale", "numero_assurance_sociale"),
    ("prenoms", "prenoms"), ("nom", "nom"), ("sexe", "sexe"),
    ("nationalite", "nationalite"), ("statut matrimonial", "statut_matrimonial"),
    ("enfants", "nombre_enfants"), ("qualification", "qualification"),
    ("niveau", "niveau_etudes"), ("poste occupe", "poste_occupe"),
    ("departement", "departement"), ("date de naissance", "date_naissance"),
    ("date d embauche", "date_embauche"), ("date de debauchage", "date_debauchage"),
    ("motif de debauchage", "motif_debauchage"), ("cnss", "date_immatriculation_cnss"),
    ("type de contrat", "type_contrat"), ("categorie", "categorie"),
    ("classe", "classe_echelon"), ("salaire de base", "salaire_base"),
    ("contact", "contact_telephonique"), ("observation", "observation"),
    ("type de travail", "type_travail"),
]
_IMPORT_DATES = {"date_naissance", "date_embauche", "date_debauchage",
                 "date_immatriculation_cnss"}


def _map_header(h):
    n = _norm(h)
    if n in ("nom", "noms"):
        return "nom"
    if "prenom" in n:
        return "prenoms"
    for key, field in _IMPORT_MAP:
        if key in n:
            return field
    return None


@frappe.whitelist()
def importer_salaries(content_base64, filename=None):
    """Importe le personnel depuis le classeur Excel de la RH (feuille « Personnel »).
    Idempotent par matricule (crée ou met à jour). Crée les départements manquants."""
    _guard()
    import base64
    import io
    import openpyxl

    raw = base64.b64decode(content_base64)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = None
    for cand in wb.worksheets:
        if "personnel" in _norm(cand.title):
            ws = cand
            break
    ws = ws or wb.worksheets[0]

    # Repérer la ligne d'en-tête (celle qui contient « Mle » ou « Nom »).
    rows = list(ws.iter_rows(values_only=True))
    hidx = None
    for i, r in enumerate(rows[:12]):
        cells = [_norm(c) for c in r]
        if "mle" in cells or "nom" in cells:
            hidx = i
            break
    if hidx is None:
        frappe.throw(_("En-tête introuvable (colonne « Mle » ou « Nom » attendue)."))

    headers = rows[hidx]
    colmap = {}
    for ci, h in enumerate(headers):
        f = _map_header(h)
        if f and f not in colmap.values():
            colmap[ci] = f

    crees, maj, ignores = 0, 0, 0
    for r in rows[hidx + 1:]:
        data = {}
        for ci, field in colmap.items():
            if ci < len(r):
                data[field] = r[ci]
        mat = data.get("matricule")
        mat = str(int(mat)) if isinstance(mat, float) else (str(mat).strip() if mat else "")
        if not mat or not (data.get("nom")):
            ignores += 1
            continue
        # Département : créer si absent.
        dep = (str(data.get("departement")).strip() if data.get("departement") else "")
        if dep and not frappe.db.exists("Departement KYA", dep):
            frappe.get_doc({"doctype": "Departement KYA", "code": dep,
                            "libelle": dep, "actif": 1}).insert(ignore_permissions=True)
        # Dates : normaliser (openpyxl renvoie déjà des datetime).
        for df in _IMPORT_DATES:
            v = data.get(df)
            if hasattr(v, "date"):
                data[df] = v.date()
            elif isinstance(v, str) and v.strip():
                data[df] = frappe.utils.getdate(v)
            else:
                data[df] = None
        if data.get("nombre_enfants") in (None, ""):
            data["nombre_enfants"] = 0
        data["matricule"] = mat

        if frappe.db.exists("Salarie KYA", mat):
            doc = frappe.get_doc("Salarie KYA", mat)
            doc.update({k: v for k, v in data.items() if k != "matricule"})
            doc.flags.ignore_permissions = True
            doc.save(ignore_permissions=True)
            maj += 1
        else:
            doc = frappe.get_doc(dict(doctype="Salarie KYA", **data))
            doc.flags.ignore_permissions = True
            doc.insert(ignore_permissions=True)
            crees += 1

    frappe.db.commit()
    out = {"crees": crees, "mis_a_jour": maj, "ignores": ignores,
           "colonnes_reconnues": len(colmap)}
    # Les autres feuilles du classeur (Stagiaires / Prestataires) sont
    # importées dans la foulée si présentes — « tout digitaliser ».
    try:
        out["stagiaires"] = _importer_feuille_annexe(
            wb, "stagiaires", "Stagiaire RH KYA", _MAP_STAGIAIRE,
            cle=("nom", "prenoms", "date_debut"), dates=("date_debut", "date_fin"),
            dep_field="departement_accueil")
        out["prestataires"] = _importer_feuille_annexe(
            wb, "prestataires", "Prestataire KYA", _MAP_PRESTATAIRE,
            cle=("raison_sociale", "date_debut_contrat"),
            dates=("date_debut_contrat", "date_fin_contrat"),
            dep_field="departement_beneficiaire")
        frappe.db.commit()
    except Exception:
        frappe.log_error(frappe.get_traceback(), "rh_effectifs.import_annexes")
    return out


# En-têtes (normalisés, sous-chaîne) des feuilles annexes du classeur.
_MAP_STAGIAIRE = [
    ("prenoms", "prenoms"), ("nom", "nom"), ("sexe", "sexe"),
    ("contact", "contact"), ("etablissement", "etablissement"),
    ("niveau", "niveau_filiere"), ("theme", "theme_stage"),
    ("departement d accueil", "departement_accueil"), ("encadreur", "encadreur"),
    ("date de debut", "date_debut"), ("date de fin", "date_fin"),
    ("convention", "convention_signee"), ("indemnite", "indemnite_stage"),
    ("observation", "observation"),
]
_MAP_PRESTATAIRE = [
    ("raison sociale", "raison_sociale"), ("nom", "raison_sociale"),
    ("type de prestataire", "type_prestataire"), ("contact", "contact"),
    ("email", "email"), ("nature de la prestation", "nature_prestation"),
    ("departement beneficiaire", "departement_beneficiaire"),
    ("date de debut", "date_debut_contrat"), ("date de fin", "date_fin_contrat"),
    ("montant", "montant_contrat"), ("mode de paiement", "mode_paiement"),
    ("reference", "reference_contrat"), ("observation", "observation"),
]


def _importer_feuille_annexe(wb, titre, doctype, mapping, cle, dates, dep_field):
    """Importe une feuille annexe (Stagiaires/Prestataires) si elle existe.
    Idempotent par `cle` (tuple de champs). Durée/statut = recalculés au save."""
    ws = None
    for cand in wb.worksheets:
        if titre in _norm(cand.title):
            ws = cand
            break
    if ws is None:
        return "feuille absente"
    rows = list(ws.iter_rows(values_only=True))
    hidx = None
    for i, r in enumerate(rows[:12]):
        cells = [_norm(c) for c in r]
        if "nom" in cells or any("raison sociale" in c for c in cells):
            hidx = i
            break
    if hidx is None:
        return "en-tête introuvable"
    colmap = {}
    for ci, h in enumerate(rows[hidx]):
        n = _norm(h)
        if not n or n in ("n ord", "duree", "statut"):
            continue  # n° d'ordre, durée et statut sont calculés chez nous
        for key, field in mapping:
            if key in n and field not in colmap.values():
                colmap[ci] = field
                break
    crees, maj, ignores = 0, 0, 0
    for r in rows[hidx + 1:]:
        data = {}
        for ci, field in colmap.items():
            if ci < len(r) and r[ci] not in (None, ""):
                data[field] = r[ci]
        if not data.get(cle[0]):
            ignores += 1
            continue
        for df in dates:
            v = data.get(df)
            if hasattr(v, "date"):
                data[df] = v.date()
            elif isinstance(v, str) and v.strip():
                data[df] = frappe.utils.getdate(v)
        dep = (str(data.get(dep_field)).strip() if data.get(dep_field) else "")
        if dep and not frappe.db.exists("Departement KYA", dep):
            frappe.get_doc({"doctype": "Departement KYA", "code": dep,
                            "libelle": dep, "actif": 1}).insert(ignore_permissions=True)
        filtres = {f: data.get(f) for f in cle if data.get(f)}
        existant = frappe.get_all(doctype, filters=filtres, pluck="name", limit=1)
        if existant:
            doc = frappe.get_doc(doctype, existant[0])
            doc.update(data)
            doc.flags.ignore_permissions = True
            doc.save(ignore_permissions=True)
            maj += 1
        else:
            doc = frappe.get_doc(dict(doctype=doctype, **data))
            doc.flags.ignore_permissions = True
            doc.insert(ignore_permissions=True)
            crees += 1
    return {"crees": crees, "mis_a_jour": maj, "ignores": ignores,
            "colonnes": len(colmap)}


@frappe.whitelist()
def liste_salaries():
    """Liste légère pour le sélecteur du parcours (matricule + nom + statut)."""
    _guard()
    rows = frappe.get_all(
        "Salarie KYA",
        fields=["name", "nom_complet", "poste_occupe", "date_debauchage", "motif_debauchage"],
        order_by="nom_complet", limit_page_length=0)
    return [{"name": r.name, "nom_complet": r.nom_complet or r.name,
             "poste": r.poste_occupe, "statut": _statut(r)} for r in rows]


@frappe.whitelist()
def parcours(salarie):
    """Parcours complet d'un salarié : infos + timeline (évolutions) + sous-listes."""
    _guard()
    s = frappe.get_doc("Salarie KYA", salarie)
    evs = frappe.get_all(
        "Evolution Carriere KYA", filters={"salarie": salarie},
        fields=["name", "nature_evolution", "date_debut", "date_fin", "duree",
                "type_contrat", "departement", "poste_occupe", "categorie",
                "classe_echelon", "salaire_base", "motif_evolution", "observation"],
        order_by="date_debut asc, creation asc", limit_page_length=0)
    for e in evs:
        e["date_debut"] = str(e["date_debut"] or "")
        e["date_fin"] = str(e["date_fin"] or "")

    def _of(nature):
        return [e for e in evs if (e.get("nature_evolution") or "") == nature]

    return {
        "salarie": {
            "matricule": s.matricule, "nom_complet": s.nom_complet,
            "poste_occupe": s.poste_occupe, "departement": s.departement,
            "sexe": s.sexe, "date_naissance": str(s.date_naissance or ""),
            "age": _age(s.date_naissance), "date_embauche": str(s.date_embauche or ""),
            "anciennete": _anciennete(s.date_embauche, s.date_debauchage),
            "type_contrat": s.type_contrat, "categorie": s.categorie,
            "classe_echelon": s.classe_echelon, "salaire_base": flt(s.salaire_base),
            "statut": _statut(s), "date_debauchage": str(s.date_debauchage or ""),
            "motif_debauchage": s.motif_debauchage, "photo": s.photo,
        },
        "evolutions": evs,
        "promotions": _of("Promotion"),
        "mutations": _of("Mutation"),
        "retrogradations": _of("Rétrogradation"),
    }


@frappe.whitelist()
def parcours_pdf(salarie):
    """Fiche PDF du parcours d'un salarié (logo embarqué → pas de fetch réseau)."""
    _guard()
    import base64
    from frappe.utils.pdf import get_pdf
    from kya_hr.kya_hr.api.stock_kya import _logo_data_uri

    d = parcours(salarie)
    s = d["salarie"]
    logo = _logo_data_uri()
    logo_img = "<img src='{0}' style='max-height:70px'>".format(logo) if logo else ""

    def esc(v):
        return frappe.utils.escape_html(str(v)) if v not in (None, "") else "—"

    couleur = {"Embauche": "#0f766e", "Promotion": "#16a34a", "Mutation": "#2563eb",
               "Renouvellement": "#0891b2", "Augmentation salariale": "#7c3aed",
               "Rétrogradation": "#dc2626", "Fin de contrat": "#64748b"}
    timeline = ""
    for e in d["evolutions"]:
        c = couleur.get(e["nature_evolution"], "#0f766e")
        timeline += (
            "<tr><td style='width:16%;color:#666'>{0}{1}</td>"
            "<td><b style='color:{2}'>{3}</b> — {4} · {5}<br>"
            "<span style='color:#888;font-size:10px'>{6}{7}</span></td></tr>").format(
            esc(frappe.utils.format_date(e["date_debut"]) if e["date_debut"] else ""),
            (" → " + frappe.utils.format_date(e["date_fin"])) if e["date_fin"] else " → (en cours)",
            c, esc(e["nature_evolution"]), esc(e["poste_occupe"]), esc(e["departement"]),
            esc(e["duree"]), (" · " + esc(e["motif_evolution"])) if e["motif_evolution"] else "")
    if not timeline:
        timeline = "<tr><td colspan='2' style='text-align:center;color:#999;padding:14px'>Aucun évènement de carrière enregistré.</td></tr>"

    html = """<html><head><meta charset='utf-8'><style>
      body{{font-family:Arial,sans-serif;font-size:11px;color:#333}}
      .hd{{display:table;width:100%}} .hd>div{{display:table-cell;vertical-align:middle}}
      .ttl{{font-size:20px;font-weight:800;color:#009B77;text-align:right}}
      .line{{height:3px;background:#009B77;margin:12px 0}}
      .sec{{background:#009B77;color:#fff;padding:6px 12px;font-weight:700;margin-top:12px}}
      .box{{border:1px solid #ddd;border-top:none;padding:10px 12px}}
      table{{width:100%;border-collapse:collapse}} td{{padding:5px 8px;border-bottom:1px solid #eee;vertical-align:top}}
      .kv{{display:table;width:100%}} .kv>div{{display:table-cell;width:33%;padding:3px 0}}
      .lbl{{color:#009B77;font-weight:700;font-size:9px;text-transform:uppercase}}
      .badge{{display:inline-block;padding:3px 12px;border-radius:12px;font-weight:700;font-size:11px}}
    </style></head><body>
      <div class='hd'><div>{logo}</div><div class='ttl'>Parcours du Salarié<br>
        <span style='font-size:11px;color:#F58220;font-style:italic'>Move beyond the sky!</span></div></div>
      <div class='line'></div>
      <div class='sec'>Informations personnelles</div>
      <div class='box'>
        <div style='font-size:15px;font-weight:800'>{nom} <span class='badge' style='background:{bgc};color:#fff'>{statut}</span></div>
        <div class='kv'>
          <div><div class='lbl'>Matricule</div>{mat}</div>
          <div><div class='lbl'>Poste actuel</div>{poste}</div>
          <div><div class='lbl'>Département</div>{dept}</div>
        </div>
        <div class='kv'>
          <div><div class='lbl'>Date d'embauche</div>{emb}</div>
          <div><div class='lbl'>Ancienneté</div>{anc} an(s)</div>
          <div><div class='lbl'>Âge</div>{age} ans</div>
        </div>
        <div class='kv'>
          <div><div class='lbl'>Type de contrat</div>{ctr}</div>
          <div><div class='lbl'>Catégorie</div>{cat}</div>
          <div><div class='lbl'>Classe &amp; échelon</div>{cls}</div>
        </div>
      </div>
      <div class='sec'>Historique de carrière ({n} évènement(s))</div>
      <div class='box'><table>{timeline}</table></div>
      <div style='margin-top:16px;text-align:center;font-size:9px;color:#999;border-top:2px solid #009B77;padding-top:8px'>
        KYA-Energy Group — fiche générée le {today}</div>
    </body></html>""".format(
        logo=logo_img, nom=esc(s["nom_complet"]),
        bgc="#16a34a" if s["statut"] == "Actif" else ("#7c3aed" if s["statut"] == "Retraité" else "#64748b"),
        statut=esc(s["statut"]), mat=esc(s["matricule"]), poste=esc(s["poste_occupe"]),
        dept=esc(s["departement"]), emb=esc(frappe.utils.format_date(s["date_embauche"]) if s["date_embauche"] else ""),
        anc=s["anciennete"], age=s["age"] if s["age"] is not None else "—",
        ctr=esc(s["type_contrat"]), cat=esc(s["categorie"]), cls=esc(s["classe_echelon"]),
        n=len(d["evolutions"]), timeline=timeline,
        today=frappe.utils.format_date(today()))

    return {"filename": "parcours-{0}.pdf".format(s["matricule"]),
            "content_base64": base64.b64encode(get_pdf(html)).decode("ascii")}


@frappe.whitelist()
def export_excel(departement=None):
    """Exporte le classeur RH : feuille Personnel (données + champs calculés),
    feuille Indicateurs, feuille Répartitions AVEC graphiques Excel natifs."""
    _guard()
    import base64
    import io
    import openpyxl
    from openpyxl.chart import BarChart, PieChart, Reference

    data = dashboard_data(departement=departement)
    k = data["kpi"]

    filters = {"departement": departement} if departement else {}
    sal = frappe.get_all(
        "Salarie KYA", filters=filters,
        fields=["matricule", "numero_assurance_sociale", "nom", "prenoms", "sexe",
                "nationalite", "statut_matrimonial", "nombre_enfants", "qualification",
                "niveau_etudes", "poste_occupe", "departement", "date_naissance",
                "date_embauche", "date_debauchage", "motif_debauchage",
                "date_immatriculation_cnss", "type_contrat", "categorie",
                "classe_echelon", "salaire_base", "type_travail",
                "contact_telephonique", "observation"],
        order_by="matricule", limit_page_length=0)

    wb = openpyxl.Workbook()

    # ── Feuille Personnel ──
    ws = wb.active
    ws.title = "Personnel"
    entetes = ["N° Ord", "Matricule", "N° Assurance Sociale", "Nom", "Prénoms", "Sexe",
               "Nationalité", "Statut matrimonial", "Nb enfants", "Qualification",
               "Niveau d'études", "Poste occupé", "Département", "Date naissance", "Âge",
               "Date embauche", "Ancienneté (ans)", "Date débauchage", "Motif débauchage",
               "Date immat. CNSS", "Type contrat", "Catégorie", "Classe & échelon",
               "Salaire de base", "Date retraite", "Statut", "Type travail",
               "Contact", "Observation"]
    ws.append(entetes)
    for i, r in enumerate(sal, start=1):
        age = _age(r.date_naissance)
        statut = _statut(r)
        retraite = frappe.utils.add_years(getdate(r.date_naissance), 60) if r.date_naissance else ""
        ws.append([
            i, r.matricule, r.numero_assurance_sociale, r.nom, r.prenoms, r.sexe,
            r.nationalite, r.statut_matrimonial, r.nombre_enfants, r.qualification,
            r.niveau_etudes, r.poste_occupe, r.departement,
            str(r.date_naissance or ""), age, str(r.date_embauche or ""),
            _anciennete(r.date_embauche, r.date_debauchage), str(r.date_debauchage or ""),
            r.motif_debauchage, str(r.date_immatriculation_cnss or ""), r.type_contrat,
            r.categorie, r.classe_echelon, flt(r.salaire_base), str(retraite), statut,
            r.type_travail, r.contact_telephonique, r.observation])

    # ── Feuille Indicateurs ──
    wi = wb.create_sheet("Indicateurs")
    wi.append(["Indicateur", "Valeur"])
    for lbl, key in [("Effectif total (actifs)", "effectif_total"), ("Hommes", "hommes"),
                     ("Femmes", "femmes"), ("Permanents (CDI)", "cdi"),
                     ("Temporaires (hors CDI)", "temporaires"),
                     ("Masse salariale mensuelle", "masse_salariale"),
                     ("Cadres (C)", "cadres"), ("Agents de maîtrise (AM)", "maitrise"),
                     ("Agents d'exécution (AE)", "execution"), ("Âge moyen", "age_moyen"),
                     ("Âge médian", "age_median"), ("Ancienneté moyenne (ans)", "anciennete_moyenne"),
                     ("Retraités", "retraites"), ("Sortis", "sortis"),
                     ("Stagiaires en cours", "stagiaires_en_cours"),
                     ("Prestataires en cours", "prestataires_en_cours")]:
        wi.append([lbl, k.get(key)])

    # ── Feuille Répartitions + graphiques natifs ──
    wr = wb.create_sheet("Répartitions")
    def _bloc(titre, dist, start_col):
        wr.cell(row=1, column=start_col, value=titre)
        wr.cell(row=2, column=start_col, value="Libellé")
        wr.cell(row=2, column=start_col + 1, value="Effectif")
        for j, (lab, val) in enumerate(zip(dist["labels"], dist["data"]), start=3):
            wr.cell(row=j, column=start_col, value=lab)
            wr.cell(row=j, column=start_col + 1, value=val)
        return len(dist["labels"])

    n_cat = _bloc("Par catégorie", data["par_categorie"], 1)
    n_dep = _bloc("Par département", data["par_departement"], 4)
    _bloc("Par type de contrat", data["par_contrat"], 7)
    _bloc("Par tranche d'âge", data["par_tranche"], 10)

    if n_cat:
        pie = PieChart(); pie.title = "Répartition par catégorie"
        pie.add_data(Reference(wr, min_col=2, min_row=2, max_row=2 + n_cat), titles_from_data=True)
        pie.set_categories(Reference(wr, min_col=1, min_row=3, max_row=2 + n_cat))
        wr.add_chart(pie, "A16")
    if n_dep:
        bar = BarChart(); bar.title = "Répartition par département"; bar.legend = None
        bar.add_data(Reference(wr, min_col=5, min_row=2, max_row=2 + n_dep), titles_from_data=True)
        bar.set_categories(Reference(wr, min_col=4, min_row=3, max_row=2 + n_dep))
        wr.add_chart(bar, "H16")

    # ── Feuilles Stagiaires & Prestataires (comme le classeur) ──
    wsg = wb.create_sheet("Stagiaires")
    wsg.append(["N° Ord", "Nom", "Prénoms", "Sexe", "Contact", "Établissement / École",
                "Niveau / Filière", "Thème du stage", "Département d'accueil",
                "Encadreur", "Date début", "Date fin", "Durée", "Convention signée",
                "Indemnité", "Statut", "Observation"])
    for i, s in enumerate(frappe.get_all(
            "Stagiaire RH KYA",
            fields=["nom", "prenoms", "sexe", "contact", "etablissement",
                    "niveau_filiere", "theme_stage", "departement_accueil", "encadreur",
                    "date_debut", "date_fin", "duree", "convention_signee",
                    "indemnite_stage", "statut", "observation"],
            order_by="date_debut desc", limit_page_length=0), start=1):
        wsg.append([i, s.nom, s.prenoms, s.sexe, s.contact, s.etablissement,
                    s.niveau_filiere, s.theme_stage, s.departement_accueil, s.encadreur,
                    str(s.date_debut or ""), str(s.date_fin or ""), s.duree,
                    s.convention_signee, flt(s.indemnite_stage), s.statut, s.observation])

    wpr = wb.create_sheet("Prestataires")
    wpr.append(["N° Ord", "Nom / Raison sociale", "Type", "Contact", "Email",
                "Nature de la prestation", "Département bénéficiaire", "Date début",
                "Date fin", "Durée", "Montant", "Mode de paiement", "Statut",
                "Référence contrat / Facture", "Observation"])
    for i, p in enumerate(frappe.get_all(
            "Prestataire KYA",
            fields=["raison_sociale", "type_prestataire", "contact", "email",
                    "nature_prestation", "departement_beneficiaire",
                    "date_debut_contrat", "date_fin_contrat", "duree",
                    "montant_contrat", "mode_paiement", "statut",
                    "reference_contrat", "observation"],
            order_by="date_debut_contrat desc", limit_page_length=0), start=1):
        wpr.append([i, p.raison_sociale, p.type_prestataire, p.contact, p.email,
                    p.nature_prestation, p.departement_beneficiaire,
                    str(p.date_debut_contrat or ""), str(p.date_fin_contrat or ""),
                    p.duree, flt(p.montant_contrat), p.mode_paiement, p.statut,
                    p.reference_contrat, p.observation])

    buf = io.BytesIO()
    wb.save(buf)
    return {"filename": "effectifs-kya-{0}.xlsx".format(today()),
            "content_base64": base64.b64encode(buf.getvalue()).decode("ascii")}


def _evolution(rows, periode_debut=None, periode_fin=None):
    """Effectif et masse salariale cumulés, mois par mois, sur des dates
    d'embauche/débauchage. Actif un mois M = embauché ≤ fin M et (pas débauché
    ou débauché > fin M)."""
    dates = [getdate(r.date_embauche) for r in rows if r.date_embauche]
    if not dates:
        return {"labels": [], "effectif": [], "masse": []}
    debut = getdate(periode_debut) if periode_debut else get_first_day(min(dates))
    fin = getdate(periode_fin) if periode_fin else getdate(today())
    debut = get_first_day(debut)

    labels, eff, masse = [], [], []
    cur = debut
    guard = 0
    while cur <= fin and guard < 600:
        month_end = getdate(frappe.utils.get_last_day(cur))
        n = 0
        m = 0.0
        for r in rows:
            de = getdate(r.date_embauche) if r.date_embauche else None
            if not de or de > month_end:
                continue
            db = getdate(r.date_debauchage) if r.date_debauchage else None
            if db and db <= month_end:
                continue
            n += 1
            m += flt(r.salaire_base)
        labels.append(cur.strftime("%Y-%m"))
        eff.append(n)
        masse.append(round(m / 1_000_000, 3))
        cur = add_months(cur, 1)
        guard += 1
    return {"labels": labels, "effectif": eff, "masse": masse}
