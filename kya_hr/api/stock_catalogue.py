"""Consultation du stock par les demandeurs + statistiques de période +
alertes réappro par e-mail.

- `catalogue` : ce que voient les CHEFS D'ÉQUIPE / le personnel avant une
  demande de sortie — désignation, catégorie, magasin, disponible. RIEN de
  sensible (ni seuils, ni classes, ni valeur). Tout utilisateur connecté.
- `stats_mouvements` / `rapport_periode_xlsx` : le rapport du responsable
  stock, filtrable semaine / mois / trimestre / année / dates libres,
  ventilé par type, catégorie, magasin et ÉQUIPE demandeuse.
- `envoyer_alertes_reappro` : mail quotidien des références en alerte
  (branché scheduler daily), paramétré dans « Paramètres Stock KYA ».
"""
from __future__ import annotations

import frappe
from frappe import _
from frappe.utils import cint, flt, today

from kya_hr.kya_hr.api.stock_kya import (
    _WH_NATIFS, _bucketize, _guard, _mag_label, _raw_sums,
    reapprovisionnement,
)


def _soldes_sans_garde(magasin=None):
    """Soldes calculés SANS le garde-fou rôles stock : le catalogue est ouvert
    à tout le personnel connecté (les demandeurs n'ont pas les rôles stock —
    ne PAS passer par soldes()/magasins() qui appellent _guard())."""
    agg = _bucketize(_raw_sums(magasin=magasin))
    return [d for d in agg.values() if abs(d["total"]) > 1e-9]


def _magasins_sans_garde():
    rows = frappe.get_all("Warehouse", filters={"disabled": 0, "is_group": 0},
                          fields=["name", "warehouse_name"], order_by="name asc")
    reels = [w for w in rows if (w.get("warehouse_name") or "").strip() not in _WH_NATIFS]
    return reels or rows


# ── Catalogue demandeurs ────────────────────────────────────────────────────
@frappe.whitelist()
def catalogue(recherche=None, categorie=None, magasin=None):
    """Stock consultable AVANT une demande de sortie. Connecté = autorisé."""
    if frappe.session.user in ("Guest", None, ""):
        frappe.throw(_("Connexion requise."), frappe.PermissionError)
    cats = {a["name"]: (a.get("categorie") or "Non classé")
            for a in frappe.get_all("Article KYA", fields=["name", "categorie"])}
    rows = []
    for d in sorted(_soldes_sans_garde(magasin=magasin),
                    key=lambda d: (d["magasin"], d["item_name"])):
        cat = cats.get(d["item"], "Non classé")
        if categorie and cat != categorie:
            continue
        if recherche and recherche.strip().lower() not in (d["item_name"] or "").lower():
            continue
        dispo = flt(d["bon_etat"])
        rows.append({
            "designation": d["item_name"], "categorie": cat,
            "magasin": d["magasin"], "magasin_label": _mag_label(d["magasin"]),
            "disponible": dispo,
            "statut": "Disponible" if dispo > 0 else "Indisponible",
        })
    rows.sort(key=lambda r: (r["statut"] != "Disponible", r["categorie"], r["designation"]))
    return {
        "articles": rows,
        "categories": sorted({r["categorie"] for r in rows}),
        "magasins": _magasins_sans_garde(),
    }


# ── Statistiques par période / équipe ───────────────────────────────────────
def _dates_periode(periode, date_debut=None, date_fin=None):
    """Résout (debut, fin) : période prédéfinie ou dates explicites."""
    from frappe.utils import add_days, get_first_day, get_last_day, getdate
    fin = getdate(date_fin) if date_fin else getdate(today())
    if date_debut:
        return getdate(date_debut), fin
    p = (periode or "mois").lower()
    if p == "semaine":
        return add_days(fin, -6), fin
    if p == "trimestre":
        return add_days(fin, -90), fin
    if p == "annee":
        return add_days(fin, -364), fin
    return get_first_day(fin), get_last_day(fin)  # mois courant


@frappe.whitelist()
def stats_mouvements(periode=None, date_debut=None, date_fin=None, magasin=None):
    """Totaux par type / magasin / catégorie / équipe demandeuse + détail."""
    _guard()
    debut, fin = _dates_periode(periode, date_debut, date_fin)
    filtres = {"date_mouvement": ["between", [str(debut), str(fin)]]}
    if magasin:
        filtres["magasin"] = magasin
    mvts = frappe.get_all(
        "Mouvement Stock KYA", filters=filtres,
        fields=["name", "date_mouvement", "type_mouvement", "item", "magasin",
                "quantite", "etat", "reference_doctype", "reference_name"],
        order_by="date_mouvement desc, creation desc")
    arts = {a["name"]: a for a in frappe.get_all(
        "Article KYA", fields=["name", "designation", "categorie"])}

    par_type, par_cat, par_mag, par_equipe = {}, {}, {}, {}
    # Équipe demandeuse : PV Sortie -> employee -> department.
    pv_names = {m["reference_name"] for m in mvts
                if m["reference_doctype"] == "PV Sortie Materiel" and m["reference_name"]}
    pv_emp, emp_dept = {}, {}
    if pv_names:
        for pv in frappe.get_all("PV Sortie Materiel",
                                 filters={"name": ["in", list(pv_names)]},
                                 fields=["name", "employee"]):
            pv_emp[pv["name"]] = pv.get("employee")
        emp_dept = {e["name"]: (e.get("department") or "Sans département")
                    for e in frappe.get_all("Employee", fields=["name", "department"])}
    detail = []
    for m in mvts:
        art = arts.get(m["item"], {})
        cat = art.get("categorie") or "Non classé"
        q = abs(flt(m["quantite"]))
        t = m["type_mouvement"]
        par_type[t] = round(par_type.get(t, 0) + q, 2)
        b = par_cat.setdefault(cat, {"categorie": cat, "entrees": 0, "sorties": 0, "retours": 0})
        g = par_mag.setdefault(m["magasin"], {"magasin": _mag_label(m["magasin"]),
                                              "entrees": 0, "sorties": 0, "retours": 0})
        key = {"Entrée": "entrees", "Sortie": "sorties", "Retour": "retours"}.get(t)
        if key:
            b[key] = round(b[key] + q, 2)
            g[key] = round(g[key] + q, 2)
        equipe = ""
        if m["reference_doctype"] == "PV Sortie Materiel":
            emp = pv_emp.get(m["reference_name"])
            equipe = emp_dept.get(emp, "Sans département") if emp else "Sans département"
            if t == "Sortie":
                e = par_equipe.setdefault(equipe, {"equipe": equipe, "sorties": 0, "mouvements": 0})
                e["sorties"] = round(e["sorties"] + q, 2)
                e["mouvements"] += 1
        detail.append({
            "date": str(m["date_mouvement"]), "type": t,
            "designation": art.get("designation") or m["item"],
            "categorie": cat, "magasin": _mag_label(m["magasin"]),
            "quantite": flt(m["quantite"]), "etat": m["etat"],
            "reference": m["reference_name"] or "", "equipe": equipe,
        })
    return {
        "periode": {"debut": str(debut), "fin": str(fin)},
        "kpi": {
            "entrees": par_type.get("Entrée", 0), "sorties": par_type.get("Sortie", 0),
            "retours": par_type.get("Retour", 0),
            "ajustements": round(sum(v for k, v in par_type.items()
                                     if k not in ("Entrée", "Sortie", "Retour")), 2),
            "nb_mouvements": len(mvts),
        },
        "par_categorie": sorted(par_cat.values(), key=lambda x: -(x["entrees"] + x["sorties"])),
        "par_magasin": sorted(par_mag.values(), key=lambda x: x["magasin"]),
        "par_equipe": sorted(par_equipe.values(), key=lambda x: -x["sorties"]),
        "detail": detail[:400],
    }


@frappe.whitelist()
def rapport_periode_xlsx(periode=None, date_debut=None, date_fin=None, magasin=None):
    """Rapport Excel de la période (synthèse + détail) pour le responsable."""
    _guard()
    import base64
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = stats_mouvements(periode, date_debut, date_fin, magasin)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0D7377")

    ws = wb.active
    ws.title = "Synthèse"
    ws.append(["Rapport de stock — période du %s au %s"
               % (data["periode"]["debut"], data["periode"]["fin"])])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Entrées", "Sorties", "Retours", "Ajustements", "Mouvements"])
    for c in ws[3]:
        c.font, c.fill = head, fill
    k = data["kpi"]
    ws.append([k["entrees"], k["sorties"], k["retours"], k["ajustements"], k["nb_mouvements"]])
    ws.append([])
    ws.append(["Par équipe demandeuse (sorties)"])
    ws.append(["Équipe", "Quantité sortie", "Nb mouvements"])
    for c in ws[7]:
        c.font, c.fill = head, fill
    for e in data["par_equipe"]:
        ws.append([e["equipe"], e["sorties"], e["mouvements"]])
    ws.append([])
    r0 = ws.max_row + 1
    ws.append(["Par catégorie"])
    ws.append(["Catégorie", "Entrées", "Sorties", "Retours"])
    for c in ws[r0 + 1]:
        c.font, c.fill = head, fill
    for b in data["par_categorie"]:
        ws.append([b["categorie"], b["entrees"], b["sorties"], b["retours"]])
    for col, w in zip("ABCDE", (34, 16, 14, 14, 14)):
        ws.column_dimensions[col].width = w

    wd = wb.create_sheet("Détail")
    wd.append(["Date", "Type", "Désignation", "Catégorie", "Magasin",
               "Quantité", "État", "Référence", "Équipe"])
    for c in wd[1]:
        c.font, c.fill = head, fill
    for m in data["detail"]:
        wd.append([m["date"], m["type"], m["designation"], m["categorie"],
                   m["magasin"], m["quantite"], m["etat"], m["reference"], m["equipe"]])
    for col, w in zip("ABCDEFGHI", (12, 10, 38, 24, 22, 10, 12, 20, 24)):
        wd.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    return {"filename": "rapport_stock_%s_%s.xlsx"
            % (data["periode"]["debut"], data["periode"]["fin"]),
            "content_base64": base64.b64encode(buf.getvalue()).decode()}


# ── Alertes e-mail réappro (scheduler daily) ────────────────────────────────
def envoyer_alertes_reappro():
    """Mail quotidien des références en alerte, selon « Paramètres Stock KYA »."""
    try:
        par = frappe.get_single("Parametres Stock KYA")
    except Exception:
        return
    if not cint(par.get("alertes_actives")):
        return
    frappe.local._kya_regles_reappro = None  # règles fraîches
    frappe.set_user("Administrator")
    rows = [r for r in reapprovisionnement(only_alertes=1)
            if r["statut"] == "RUPTURE" or cint(par.get("inclure_a_commander"))]
    if not rows:
        return
    dests = [l.strip() for l in (par.get("destinataires") or "").splitlines() if l.strip()]
    if not dests:
        dests = [u["email"] for u in frappe.get_all(
            "Has Role", filters={"role": "Responsable Stock", "parenttype": "User"},
            fields=["parent as email"])
            if frappe.db.get_value("User", u["email"], "enabled")]
    if not dests:
        return
    lignes = "".join(
        "<tr><td>{d}</td><td>{m}</td><td style='text-align:center'>{b}</td>"
        "<td style='text-align:center'>{s}</td><td style='text-align:center;"
        "font-weight:700;color:{c}'>{st}</td><td style='text-align:center'>{q}</td></tr>"
        .format(d=frappe.utils.escape_html(r["designation"]),
                m=frappe.utils.escape_html(r["magasin_label"]),
                b=r["dispo"], s=r["seuil_mini"],
                c="#dc2626" if r["statut"] == "RUPTURE" else "#d97706",
                st=r["statut"], q=r["qte_a_commander"] or "-")
        for r in rows)
    html = ("<p>Bonjour,</p><p>Références en alerte de réapprovisionnement ce jour :</p>"
            "<table border='1' cellpadding='6' cellspacing='0' "
            "style='border-collapse:collapse;font-size:13px'>"
            "<tr style='background:#0d7377;color:#fff'><th>Désignation</th><th>Magasin</th>"
            "<th>Dispo (bon état)</th><th>Seuil</th><th>Statut</th><th>À commander</th></tr>"
            + lignes + "</table>"
            "<p>Détail et actions : <a href='/stock-kya'>cockpit Stock KYA</a> "
            "(onglet Réappro &amp; Alertes).</p><p>— ERP KYA</p>")
    frappe.sendmail(recipients=dests,
                    subject="[Stock KYA] %d référence(s) en alerte de réappro" % len(rows),
                    message=html)
