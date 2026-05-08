import re
from datetime import datetime

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import flt, now_datetime
from frappe.utils.file_manager import get_file_path


_ALLOWED_ROLES = {"Comptable", "DFC", "DAAF", "Accounts Manager", "System Manager"}


class KYAComptaImport(Document):
    def validate(self):
        self._set_trace_fields()
        if self.source_file and self.has_value_changed("source_file"):
            self.reimport_from_source()

    def _set_trace_fields(self):
        if not self.imported_by:
            self.imported_by = frappe.session.user
        if not self.date_import:
            self.date_import = now_datetime()
        if not self.statut_import:
            self.statut_import = "Nouveau"

    @frappe.whitelist()
    def reimport_from_source(self):
        self._check_permission()
        if not self.source_file:
            frappe.throw(_("Veuillez joindre le fichier Excel à importer."))

        try:
            rows = parse_compta_excel(self.source_file, self.type_document)
            self.set("lignes", [])
            totals = {
                "total_lignes": 0,
                "total_debit": 0,
                "total_credit": 0,
                "total_salaire_net": 0,
                "total_facture": 0,
            }
            for row in rows:
                self.append("lignes", row)
                totals["total_lignes"] += 1
                totals["total_debit"] += flt(row.get("debit"))
                totals["total_credit"] += flt(row.get("credit"))
                totals["total_salaire_net"] += flt(row.get("net_a_payer"))
                if row.get("type_ligne") == "Facture":
                    totals["total_facture"] += flt(row.get("montant"))

            for field, value in totals.items():
                self.set(field, value)
            self.statut_import = "Importé"
            self.message_erreur = None
            if not self.periode:
                self.periode = infer_period(self.type_document, rows) or self.periode
        except Exception as exc:
            self.statut_import = "Erreur"
            self.message_erreur = str(exc)[:1000]
            frappe.log_error(frappe.get_traceback(), "KYA Compta Import")
            raise

    def _check_permission(self):
        roles = set(frappe.get_roles(frappe.session.user))
        if not roles.intersection(_ALLOWED_ROLES):
            frappe.throw(_("Accès réservé à la comptabilité."), frappe.PermissionError)


def parse_compta_excel(file_url, type_document):
    try:
        from openpyxl import load_workbook
    except ImportError:
        frappe.throw(_("La librairie openpyxl est requise pour lire les fichiers Excel."))

    file_path = get_file_path(file_url)
    workbook = load_workbook(file_path, data_only=True, read_only=True)

    if type_document == "Etat de Salaire":
        return parse_salary(workbook)
    if type_document == "Grand Livre Général":
        return parse_general_ledger(workbook)
    if type_document == "Facture":
        return parse_invoice(workbook)

    frappe.throw(_("Type de document non pris en charge: {0}").format(type_document))


def parse_salary(workbook):
    sheet = workbook["JANVIER"] if "JANVIER" in workbook.sheetnames else workbook.worksheets[0]
    rows = []
    for row_index in range(6, sheet.max_row + 1):
        number = cell_value(sheet, row_index, 1)
        if is_total_marker(number):
            break
        matricule = cell_value(sheet, row_index, 6)
        employe = cell_value(sheet, row_index, 5)
        if not matricule and not employe:
            continue
        rows.append({
            "type_ligne": "Salaire",
            "row_index": row_index,
            "matricule": as_text(matricule),
            "employe": as_text(employe),
            "situation_matrimoniale": as_text(cell_value(sheet, row_index, 2)),
            "profession": as_text(cell_value(sheet, row_index, 8)),
            "salaire_base": flt(cell_value(sheet, row_index, 15)),
            "salaire_brut": flt(cell_value(sheet, row_index, 24)),
            "total_retenues": flt(cell_value(sheet, row_index, 36)),
            "net_a_payer": flt(cell_value(sheet, row_index, 52)),
            "charges_patronales": flt(cell_value(sheet, row_index, 53)),
            "libelle": "Etat de salaire {}".format(as_text(cell_value(sheet, row_index, 14))),
            "raw_json": frappe.as_json(row_to_dict(sheet, row_index, 1, 56)),
        })
    return rows


def parse_general_ledger(workbook):
    sheet = workbook["Grand-livre des comptes"] if "Grand-livre des comptes" in workbook.sheetnames else workbook.worksheets[0]
    rows = []
    current_account = None
    for row_index in range(1, sheet.max_row + 1):
        first = cell_value(sheet, row_index, 1)
        third = cell_value(sheet, row_index, 3)
        if is_account_row(first, third):
            current_account = as_text(first)
            rows.append({
                "type_ligne": "Compte",
                "row_index": row_index,
                "compte": current_account,
                "libelle": as_text(third),
                "raw_json": frappe.as_json(row_to_dict(sheet, row_index, 1, sheet.max_column)),
            })
            continue
        if not looks_like_date(first):
            continue
        rows.append({
            "type_ligne": "Grand Livre",
            "row_index": row_index,
            "date_operation": normalize_date(first),
            "compte": current_account,
            "code_journal": as_text(cell_value(sheet, row_index, 2)),
            "numero_piece": as_text(third),
            "libelle": as_text(cell_value(sheet, row_index, 5) or cell_value(sheet, row_index, 6)),
            "debit": flt(cell_value(sheet, row_index, 9) or cell_value(sheet, row_index, 10)),
            "credit": flt(cell_value(sheet, row_index, 12) or cell_value(sheet, row_index, 13)),
            "solde_progressif": flt(cell_value(sheet, row_index, 14) or cell_value(sheet, row_index, 15)),
            "raw_json": frappe.as_json(row_to_dict(sheet, row_index, 1, sheet.max_column)),
        })
    return rows


def parse_invoice(workbook):
    rows = []
    for sheet in workbook.worksheets:
        numero = extract_invoice_number(sheet)
        objet = extract_invoice_object(sheet)
        for row_index in range(1, sheet.max_row + 1):
            designation = cell_value(sheet, row_index, 3)
            if not designation:
                continue
            line_no = cell_value(sheet, row_index, 2)
            montant = flt(cell_value(sheet, row_index, 7))
            if not (is_number_like(line_no) and (montant or cell_value(sheet, row_index, 5))):
                continue
            rows.append({
                "type_ligne": "Facture",
                "row_index": row_index,
                "numero_facture": numero,
                "objet_facture": objet,
                "designation": as_text(designation),
                "unite": as_text(cell_value(sheet, row_index, 4)),
                "quantite": flt(cell_value(sheet, row_index, 5)),
                "prix_unitaire": flt(cell_value(sheet, row_index, 6)),
                "montant": montant,
                "libelle": as_text(designation),
                "raw_json": frappe.as_json(row_to_dict(sheet, row_index, 1, sheet.max_column)),
            })
    return rows


def infer_period(type_document, rows):
    if type_document == "Facture" and rows:
        return rows[0].get("numero_facture")
    if type_document == "Etat de Salaire" and rows:
        match = re.search(r"(\d{2}/\d{2}/\d{4}.*)", rows[0].get("libelle") or "")
        return match.group(1) if match else None
    return None


def extract_invoice_number(sheet):
    for row_index in range(1, min(sheet.max_row, 12) + 1):
        value = as_text(cell_value(sheet, row_index, 2))
        if value.startswith("N°"):
            return value.replace("N°", "").strip()
    return None


def extract_invoice_object(sheet):
    for row_index in range(1, min(sheet.max_row, 14) + 1):
        value = as_text(cell_value(sheet, row_index, 2))
        if value.lower().startswith("objet"):
            return value.split(":", 1)[-1].strip()
    return None


def cell_value(sheet, row, column):
    return sheet.cell(row=row, column=column).value


def row_to_dict(sheet, row_index, start_col, end_col):
    return {
        str(col): normalize_raw_value(cell_value(sheet, row_index, col))
        for col in range(start_col, end_col + 1)
        if cell_value(sheet, row_index, col) not in (None, "")
    }


def normalize_raw_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def looks_like_date(value):
    return isinstance(value, datetime) or hasattr(value, "year")


def is_account_row(first, third):
    first_text = as_text(first)
    return first_text.isdigit() and len(first_text) >= 3 and bool(third) and not looks_like_date(first)


def is_total_marker(value):
    return as_text(value).upper().startswith("TOTAL")


def is_number_like(value):
    if isinstance(value, (int, float)):
        return True
    return as_text(value).isdigit()


def as_text(value):
    if value is None:
        return ""
    return str(value).strip()
